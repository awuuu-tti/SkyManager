from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
import shutil

from filelock import FileLock
from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / 'data'
BACKUP_DIR = BASE_DIR / 'backups'
DB_PATH = DATA_DIR / 'airline_data.xlsx'
LOCK_PATH = DATA_DIR / 'airline_data.lock'

HEADERS = {
    'users': ['id', 'full_name', 'email', 'phone', 'birth_date', 'password_hash', 'role', 'is_active', 'loyalty_member', 'created_at'],
    'planes': ['id', 'model', 'tail_number', 'seats'],
    'flights': ['id', 'flight_number', 'origin', 'destination', 'departure_date', 'departure_time', 'arrival_time', 'plane_id', 'plane_model', 'total_seats', 'free_seats', 'base_price', 'status', 'fare_type', 'created_at'],
    'bookings': ['id', 'booking_code', 'user_id', 'flight_id', 'flight_number', 'passenger_name', 'passenger_category', 'discount_percent', 'final_price', 'status', 'booked_at', 'cancel_allowed', 'cancelled_at'],
    'discount_rules': ['category', 'label', 'discount_percent'],
    'logs': ['id', 'event_time', 'actor_email', 'actor_role', 'action', 'details'],
}

DISCOUNTS = {'adult': 0, 'child': 50, 'loyal': 15, 'employee': 30}
LABELS = {'adult': 'Взрослый', 'child': 'Ребенок', 'loyal': 'Постоянный клиент', 'employee': 'Сотрудник'}
ROLES = {'client': 'Клиент', 'employee': 'Сотрудник', 'admin': 'Администратор'}
STATUSES = ['По расписанию', 'Задерживается', 'Отменен']


def now():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def ensure_db():
    DATA_DIR.mkdir(exist_ok=True)
    BACKUP_DIR.mkdir(exist_ok=True)
    if DB_PATH.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = 'users'
    for name in HEADERS:
        if name not in wb.sheetnames:
            wb.create_sheet(name)
        sheet = wb[name]
        sheet.append(HEADERS[name])
    for row in [
        [1, 'Администратор SkyManager', 'admin@skymanager.local', '+79990000001', '1990-01-01', generate_password_hash('Admin123!'), 'admin', 1, 1, now()],
        [2, 'Мария Диспетчер', 'employee@skymanager.local', '+79990000002', '1992-05-16', generate_password_hash('Employee123!'), 'employee', 1, 1, now()],
        [3, 'Иван Пассажир', 'client@skymanager.local', '+79990000003', '1998-07-11', generate_password_hash('Client123!'), 'client', 1, 1, now()],
    ]:
        wb['users'].append(row)
    for row in [[1, 'Airbus A320', 'RA-32001', 180], [2, 'Boeing 737-800', 'RA-73802', 189], [3, 'Sukhoi Superjet 100', 'RA-10003', 98]]:
        wb['planes'].append(row)
    for row in [
        [1, 'SM-101', 'Москва', 'Сочи', '2026-04-20', '09:30', '13:05', 1, 'Airbus A320', 180, 142, 12500, 'По расписанию', 'Standard', now()],
        [2, 'SM-202', 'Санкт-Петербург', 'Казань', '2026-04-21', '14:10', '16:20', 3, 'Sukhoi Superjet 100', 98, 64, 7900, 'По расписанию', 'Flexible', now()],
        [3, 'SM-303', 'Екатеринбург', 'Новосибирск', '2026-04-22', '07:45', '10:20', 2, 'Boeing 737-800', 189, 121, 9800, 'Задерживается', 'Promo', now()],
        [4, 'SM-404', 'Москва', 'Калининград', '2026-04-23', '18:00', '20:15', 3, 'Sukhoi Superjet 100', 98, 47, 8900, 'По расписанию', 'Standard', now()],
    ]:
        wb['flights'].append(row)
    for row in [[k, LABELS[k], v] for k, v in DISCOUNTS.items()]:
        wb['discount_rules'].append(row)
    wb['logs'].append([1, now(), 'system', 'system', 'init_db', 'Создана демонстрационная база'])
    wb.save(DB_PATH)


def read_sheet(name: str):
    ensure_db()
    wb = load_workbook(DB_PATH)
    ws = wb[name]
    headers = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if any(v is not None for v in row.values()):
            rows.append(row)
    return rows


def write_sheet(name: str, rows: list[dict]):
    ensure_db()
    with FileLock(str(LOCK_PATH), timeout=10):
        wb = load_workbook(DB_PATH)
        ws = wb[name]
        ws.delete_rows(2, ws.max_row)
        for row in rows:
            ws.append([row.get(h) for h in HEADERS[name]])
        wb.save(DB_PATH)


def next_id(name: str):
    rows = read_sheet(name)
    return max([int(r.get('id') or 0) for r in rows] + [0]) + 1


def add_row(name: str, row: dict):
    rows = read_sheet(name)
    rows.append(row)
    write_sheet(name, rows)


def user_by_email(email: str):
    email = email.lower().strip()
    return next((u for u in read_sheet('users') if str(u['email']).lower() == email and int(u['is_active'] or 0) == 1), None)


def user_by_id(user_id: int):
    return next((u for u in read_sheet('users') if int(u['id']) == int(user_id) and int(u['is_active'] or 0) == 1), None)


def flight_by_id(flight_id: int):
    return next((f for f in read_sheet('flights') if int(f['id']) == int(flight_id)), None)


def booking_by_id(booking_id: int):
    return next((b for b in read_sheet('bookings') if int(b['id']) == int(booking_id)), None)


def active_bookings_for_flight(flight_id: int):
    users = {int(u['id']): u for u in read_sheet('users') if int(u.get('is_active') or 0) == 1}
    rows = [b for b in read_sheet('bookings') if int(b['flight_id']) == int(flight_id) and b['status'] == 'active']
    for row in rows:
        user = users.get(int(row['user_id'])) or {}
        row['user_phone'] = user.get('phone')
        row['user_email'] = user.get('email')
    return rows


def save_user(row: dict):
    users = read_sheet('users')
    for i, user in enumerate(users):
        if int(user['id']) == int(row['id']):
            users[i] = row
            break
    else:
        users.append(row)
    write_sheet('users', users)


def save_flight(row: dict):
    flights = read_sheet('flights')
    for i, flight in enumerate(flights):
        if int(flight['id']) == int(row['id']):
            flights[i] = row
            break
    else:
        flights.append(row)
    write_sheet('flights', flights)


def remove_flight(flight_id: int):
    write_sheet('flights', [f for f in read_sheet('flights') if int(f['id']) != int(flight_id)])


def remove_user(user_id: int):
    users = read_sheet('users')
    for user in users:
        if int(user['id']) == int(user_id):
            user['is_active'] = 0
    write_sheet('users', users)


def reset_password(user_id: int, password: str):
    users = read_sheet('users')
    for user in users:
        if int(user['id']) == int(user_id):
            user['password_hash'] = generate_password_hash(password)
    write_sheet('users', users)


def add_booking(user: dict, flight: dict, passengers: list[dict]):
    bookings = read_sheet('bookings')
    flights = read_sheet('flights')
    current = next((f for f in flights if int(f['id']) == int(flight['id'])), None)
    if not current or int(current['free_seats']) < len(passengers):
        return False
    current['free_seats'] = int(current['free_seats']) - len(passengers)
    code_base = max([int(str(b.get('booking_code', 'BK-100000')).split('-')[-1]) for b in bookings] + [100000])
    for n, passenger in enumerate(passengers, start=1):
        bookings.append({
            'id': next_id('bookings') + n - 1,
            'booking_code': f'BK-{code_base + n}',
            'user_id': int(user['id']),
            'flight_id': int(current['id']),
            'flight_number': current['flight_number'],
            'passenger_name': passenger['name'],
            'passenger_category': passenger['category'],
            'discount_percent': DISCOUNTS[passenger['category']],
            'final_price': int(float(current['base_price']) * (100 - DISCOUNTS[passenger['category']]) / 100),
            'status': 'active',
            'booked_at': now(),
            'cancel_allowed': 0 if current['fare_type'] == 'Promo' else 1,
            'cancelled_at': None,
        })
    write_sheet('bookings', bookings)
    write_sheet('flights', flights)
    return True


def cancel_booking(booking_id: int):
    bookings = read_sheet('bookings')
    flights = read_sheet('flights')
    target = next((b for b in bookings if int(b['id']) == int(booking_id)), None)
    if not target or target['status'] != 'active':
        return
    target['status'] = 'cancelled'
    target['cancelled_at'] = now()
    for flight in flights:
        if int(flight['id']) == int(target['flight_id']):
            flight['free_seats'] = int(flight['free_seats']) + 1
    write_sheet('bookings', bookings)
    write_sheet('flights', flights)


def bookings_for_user(user_id: int):
    flights = {int(f['id']): deepcopy(f) for f in read_sheet('flights')}
    rows = [b for b in read_sheet('bookings') if int(b['user_id']) == int(user_id)]
    for row in rows:
        row['flight'] = flights.get(int(row['flight_id']))
    return rows


def report_rows(start_date: str | None, end_date: str | None):
    rows = [b for b in read_sheet('bookings') if b['status'] == 'active']
    if start_date:
        rows = [b for b in rows if str(b['booked_at'])[:10] >= start_date]
    if end_date:
        rows = [b for b in rows if str(b['booked_at'])[:10] <= end_date]
    return rows


def backup_database():
    ensure_db()
    name = BACKUP_DIR / f"airline_data_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    with FileLock(str(LOCK_PATH), timeout=10):
        shutil.copy2(DB_PATH, name)
    return name


def add_log(action: str, details: str, user: dict | None = None):
    logs = read_sheet('logs')
    logs.append({'id': next_id('logs'), 'event_time': now(), 'actor_email': (user or {}).get('email', 'guest'), 'actor_role': (user or {}).get('role', 'guest'), 'action': action, 'details': details})
    write_sheet('logs', logs)
